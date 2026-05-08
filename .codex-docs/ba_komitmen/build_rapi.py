from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from table_geometry import apply_table_geometry, exact_column_widths, section_content_width_dxa


BASE_DIR = Path("/Users/robbysyamsuddin/works/inventoryapps/.codex-docs/ba_komitmen")
OUTPUT = BASE_DIR / "BA_Komitmen_Harian_PV_rapi.docx"

HEADER_FILL = "DCE6F1"
HEADER_TEXT = RGBColor(31, 78, 121)
BORDER_COLOR = "B7C9E2"
TEXT_COLOR = RGBColor(34, 34, 34)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER_COLOR, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        elem = tc_borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            tc_borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), size)
        elem.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_paragraph(paragraph, *, align=None, bold=False, size=11, color=TEXT_COLOR, space_after=4):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def format_cell_text(cell, *, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size=9, color=TEXT_COLOR):
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def configure_section(section, *, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21)
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.2)
        section.right_margin = Cm(1.2)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = "Berita Acara Komitmen Harian Plan Inject PV"
    set_paragraph(header, align=WD_ALIGN_PARAGRAPH.RIGHT, size=8, color=RGBColor(99, 99, 99), space_after=0)

    footer = section.footer.paragraphs[0]
    footer.text = ""


def add_commitment_table(document, title_text):
    heading = document.add_paragraph()
    heading.style = document.styles["Heading 1"]
    heading_run = heading.add_run(title_text)
    heading_run.font.name = "Arial"
    heading_run.font.size = Pt(12)
    heading_run.font.bold = True
    heading_run.font.color.rgb = HEADER_TEXT
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(6)

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    headers = [
        "Tanggal",
        "Redeem PV Reguler",
        "Plan Inject PV Reguler",
        "Redeem PV By.U",
        "Plan Inject PV By.U",
    ]
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, text in enumerate(headers):
        cell = header_row.cells[idx]
        cell.text = text
        set_cell_shading(cell, HEADER_FILL)
        set_cell_border(cell, color="9FBAD0", size="8")
        format_cell_text(cell, bold=True, size=8, color=HEADER_TEXT)

    for day in range(1, 32):
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Cm(0.5)
        row.cells[0].text = str(day)
        for idx, cell in enumerate(row.cells):
            set_cell_border(cell)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            format_cell_text(cell, align=align, size=8)

    widths = exact_column_widths([1100, 2200, 2500, 1700, 1860], section_content_width_dxa(document.sections[-1]))
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=sum(widths),
        indent_dxa=0,
        cell_margins_dxa={"top": 35, "bottom": 35, "start": 80, "end": 80},
    )
    document.add_paragraph()


def main():
    doc = Document()

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT_COLOR

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(15)
    title_style.font.bold = True
    title_style.font.color.rgb = HEADER_TEXT

    heading_style = styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True
    heading_style.font.color.rgb = HEADER_TEXT

    configure_section(doc.sections[0], landscape=False)
    for section in doc.sections:
        add_header_footer(section)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("BERITA ACARA")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("KOMITMEN HARIAN PLAN INJECT PV REGULER DAN PV BY.U DAILY")
    set_paragraph(subtitle, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, color=TEXT_COLOR, space_after=12)

    nomor = doc.add_paragraph()
    nomor.add_run("Nomor: ________________________________")
    set_paragraph(nomor, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=10)

    intro = doc.add_paragraph()
    intro.add_run(
        "Pada hari ini, __________ tanggal ___ bulan __________ tahun 20___, "
        "kami menyatakan komitmen terhadap target harian Plan Inject PV Reguler "
        "dan PV By.U Daily sebagai berikut:"
    )
    intro.paragraph_format.first_line_indent = Cm(0.75)
    set_paragraph(intro, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=12)

    landscape_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape_section, landscape=True)
    add_header_footer(landscape_section)

    for idx, label in enumerate(("1D", "3D", "5D", "7D", "14D", "28D")):
        add_commitment_table(doc, f"Plan Inject PV {label} (Per Tanggal)")
        if idx != 5:
            doc.add_page_break()

    closing_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(closing_section, landscape=False)
    add_header_footer(closing_section)

    notes_head = doc.add_paragraph()
    notes_head.add_run("Catatan")
    set_paragraph(notes_head, bold=True, size=11, color=HEADER_TEXT, space_after=4)

    notes = [
        "Target disusun berdasarkan data redeem harian.",
        "Target mempertimbangkan estimasi stock days di warehouse dan outlet.",
        "Monitoring dilakukan setiap hari dan direkap secara bulanan.",
    ]
    for note in notes:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(note)
        set_paragraph(p, size=10, space_after=2)

    doc.add_paragraph()

    penutup_head = doc.add_paragraph(style="Heading 1")
    penutup_head.add_run("Penutup")
    set_paragraph(penutup_head, bold=True, size=12, color=HEADER_TEXT, space_after=4)

    closing = doc.add_paragraph()
    closing.add_run(
        "Demikian Berita Acara ini dibuat sebagai bentuk komitmen bersama "
        "dalam mencapai target harian yang telah ditetapkan."
    )
    set_paragraph(closing, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=14)

    signature_intro = doc.add_paragraph()
    signature_intro.add_run("Dokumen ini disepakati dan ditandatangani oleh:")
    set_paragraph(signature_intro, size=11, space_after=6)

    sign_table = doc.add_table(rows=2, cols=4)
    sign_table.style = "Table Grid"
    sign_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    sign_headers = ["Manager Support", "Manager Cluster", "GM Cluster", "MCOT"]
    for idx, header_text in enumerate(sign_headers):
        sign_table.cell(0, idx).text = header_text
        set_cell_shading(sign_table.cell(0, idx), HEADER_FILL)
        set_cell_border(sign_table.cell(0, idx), color="9FBAD0", size="8")
        format_cell_text(sign_table.cell(0, idx), bold=True, size=9, color=HEADER_TEXT)

        sign_table.cell(1, idx).text = "\n\n(____________________________)"
        set_cell_border(sign_table.cell(1, idx))
        format_cell_text(sign_table.cell(1, idx), align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
        sign_table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        sign_table.rows[1].height = Cm(2.5)

    sign_widths = exact_column_widths([2340, 2340, 2340, 2340], section_content_width_dxa(doc.sections[-1]))
    apply_table_geometry(
        sign_table,
        sign_widths,
        table_width_dxa=sum(sign_widths),
        indent_dxa=0,
        cell_margins_dxa={"top": 120, "bottom": 120, "start": 120, "end": 120},
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
